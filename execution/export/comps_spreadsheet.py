"""Generate spreadsheet with all comparable property data."""
from datetime import datetime
from pathlib import Path
from typing import List, Dict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from loguru import logger


class CompsSpreadsheetGenerator:
    """Generate spreadsheet showing all comparable properties used in analysis."""

    def __init__(self, output_dir: str = ".tmp"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.logger = logger.bind(generator='CompsSpreadsheet')

    def generate(self, all_comps: List[Dict], location: str, filename: str = None) -> str:
        """
        Generate spreadsheet with all comparable properties.

        Args:
            all_comps: List of all comp dicts
            location: Location string (e.g., "Blue Earth, MN")
            filename: Optional custom filename

        Returns:
            Path to generated XLSX file
        """
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_location = location.replace(", ", "_").replace(" ", "_").lower()
            filename = f"comparable_properties_{safe_location}_{timestamp}.xlsx"

        filepath = self.output_dir / filename

        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Comparable Properties"

        # Define columns
        columns = [
            ("Address", 45),
            ("City", 15),
            ("State", 8),
            ("Zip Code", 10),
            ("Bedrooms", 10),
            ("Bathrooms", 10),
            ("Sqft", 10),
            ("Year Built", 12),
            ("Sale Price", 15),
            ("Price/Sqft", 12),
            ("Sale Date", 15),
            ("Distance (mi)", 12),
            ("Source", 20),
        ]

        # Write header
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)

        for col_idx, (col_name, col_width) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(col_idx)].width = col_width

        # Write data rows
        row_idx = 2
        for comp in all_comps:
            # Parse address
            address_str = comp.get('address', 'Unknown')
            if isinstance(address_str, dict):
                street = address_str.get('streetAddress', '')
                city = address_str.get('city', '')
                state = address_str.get('state', '')
                zipcode = address_str.get('zipcode', '')
            else:
                # Try to parse string address
                parts = address_str.split(', ')
                street = parts[0] if len(parts) > 0 else ''
                city = parts[1] if len(parts) > 1 else ''
                state = parts[2].split()[0] if len(parts) > 2 else ''
                zipcode = parts[2].split()[-1] if len(parts) > 2 and len(parts[2].split()) > 1 else ''

            # Address with Google Maps link
            full_address = address_str if isinstance(address_str, str) else f"{street}, {city}, {state} {zipcode}"
            maps_url = f"https://www.google.com/maps/search/{full_address.replace(' ', '+')}"
            addr_cell = ws.cell(row=row_idx, column=1, value=full_address)
            addr_cell.hyperlink = maps_url
            addr_cell.font = Font(color="0563C1", underline="single")

            ws.cell(row=row_idx, column=2, value=city)
            ws.cell(row=row_idx, column=3, value=state)
            ws.cell(row=row_idx, column=4, value=zipcode)

            # Property details
            beds = comp.get('bedrooms')
            baths = comp.get('bathrooms')
            sqft = comp.get('sqft')
            year_built = comp.get('year_built')

            ws.cell(row=row_idx, column=5, value=beds)
            ws.cell(row=row_idx, column=6, value=baths)
            ws.cell(row=row_idx, column=7, value=sqft)
            ws.cell(row=row_idx, column=8, value=year_built)

            # Sale price
            sale_price = comp.get('sale_price')
            if sale_price:
                price_cell = ws.cell(row=row_idx, column=9, value=sale_price)
                price_cell.number_format = '$#,##0'

            # Price per sqft
            if sale_price and sqft and sqft > 0:
                psf = sale_price / sqft
                psf_cell = ws.cell(row=row_idx, column=10, value=psf)
                psf_cell.number_format = '$#,##0.00'

            # Sale date
            ws.cell(row=row_idx, column=11, value=comp.get('sale_date'))

            # Distance
            distance = comp.get('distance_miles')
            if distance:
                ws.cell(row=row_idx, column=12, value=distance)

            # Source
            source = comp.get('source', 'Unknown')
            ws.cell(row=row_idx, column=13, value=source)

            row_idx += 1

        # Add summary at top
        ws.insert_rows(1, 3)

        # Title
        title_cell = ws.cell(row=1, column=1, value=f"Comparable Properties - {location}")
        title_cell.font = Font(size=16, bold=True, color="366092")

        # Summary
        summary_text = f"Total Comps: {len(all_comps)} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        summary_cell = ws.cell(row=2, column=1, value=summary_text)
        summary_cell.font = Font(size=10, italic=True)

        # Adjust header row after inserting rows
        header_row = 4
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Freeze panes at header
        ws.freeze_panes = f"A{header_row + 1}"

        # Add filters
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{row_idx - 1}"

        # Save workbook
        wb.save(filepath)

        self.logger.info(f"Comps spreadsheet generated: {filepath}")
        return str(filepath)
